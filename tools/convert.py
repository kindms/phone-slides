"""把 PPT 轉成手機簡報：每張投影片匯出成圖片＋備忘稿，產生講者／聽眾連結。

用法：py -3.14 tools/convert.py 簡報.pptx [--title 顯示標題]
需要：Windows 安裝 PowerPoint、pip install python-pptx pillow
"""
import argparse
import json
import os
import secrets
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from PIL import Image
from pptx import Presentation

for s in (sys.stdout, sys.stderr):
    s.reconfigure(encoding="utf-8")

ROOT = Path(__file__).resolve().parent.parent
SITE = "https://kindms.github.io/phone-slides/"
WIDTH = 1920

EXPORT_PS1 = r"""
$ErrorActionPreference = 'Stop'
$app = New-Object -ComObject PowerPoint.Application
$pres = $app.Presentations.Open($env:PS_SRC, $true, $false, $false)
$h = [int]([int]$env:PS_W * $pres.PageSetup.SlideHeight / $pres.PageSetup.SlideWidth)
$i = 0
foreach ($s in $pres.Slides) {
  $i++
  if ($s.SlideShowTransition.Hidden) { continue }
  $s.Export((Join-Path $env:PS_OUT ('{0:D3}.png' -f $i)), 'PNG', [int]$env:PS_W, $h)
}
$pres.Close()
if ($app.Presentations.Count -eq 0) { $app.Quit() }
"""


def export_png(src: Path, out: Path):
    ps1 = out / "_export.ps1"
    ps1.write_text(EXPORT_PS1, encoding="utf-8-sig")
    env = dict(os.environ, PS_SRC=str(src), PS_OUT=str(out), PS_W=str(WIDTH))
    subprocess.run(["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", str(ps1)],
                   env=env, check=True)


def slide_notes(src: Path):
    notes = {}
    for i, slide in enumerate(Presentation(str(src)).slides, 1):
        if slide.has_notes_slide:
            notes[i] = slide.notes_slide.notes_text_frame.text.strip()
    return notes


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("pptx")
    ap.add_argument("--title")
    a = ap.parse_args()

    src = Path(a.pptx).resolve()
    title = a.title or src.stem
    deck = secrets.token_hex(8)
    room = "".join(secrets.choice("ABCDEFGHJKLMNPQRSTUVWXYZ23456789") for _ in range(6))
    dest = ROOT / "d" / deck
    dest.mkdir(parents=True)

    print(f"轉檔中：{src.name}（PowerPoint 會在背景開啟，請稍候）")
    tmp = Path(tempfile.mkdtemp())
    slides = []
    try:
        export_png(src, tmp)
        notes = slide_notes(src)
        for png in sorted(tmp.glob("*.png")):
            idx = int(png.stem)
            name = secrets.token_hex(6) + ".jpg"   # 隨機檔名，聽眾猜不到下一張
            Image.open(png).convert("RGB").save(dest / name, "JPEG", quality=82, optimize=True, progressive=True)
            slides.append({"img": name, "notes": notes.get(idx, "")})
            print(f"  第 {idx} 張完成")
    except Exception:
        shutil.rmtree(dest, ignore_errors=True)   # 失敗就不留半成品
        raise
    finally:
        shutil.rmtree(tmp, ignore_errors=True)

    if not slides:
        shutil.rmtree(dest)
        sys.exit("沒有匯出任何投影片，請確認檔案。")

    (dest / "manifest.json").write_text(
        json.dumps({"title": title, "room": room, "slides": slides}, ensure_ascii=False, indent=1),
        encoding="utf-8")

    host = f"{SITE}host.html?d={deck}"
    audience = f"{SITE}?r={room}"
    record = f"【{title}】{len(slides)} 張\n講者遙控（只給自己）：{host}\n聽眾連結（做成 QR Code）：{audience}\n\n"
    with open(ROOT / "簡報連結.txt", "a", encoding="utf-8") as f:
        f.write(record)
    print("\n" + record + "（已記錄在 簡報連結.txt）")


if __name__ == "__main__":
    main()
